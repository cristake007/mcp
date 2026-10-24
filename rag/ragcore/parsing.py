"""Loss-aware local document parsing. Original files are never modified.

Text extracted here is preserved verbatim in PostgreSQL, but PDF layout,
spreadsheet presentation and DOCX drawing shapes cannot be proved lossless.
Each extraction unit carries a source location and warnings are surfaced.
"""
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".markdown"}


@dataclass(frozen=True)
class Unit:
    start: int
    end: int
    source: dict[str, Any]


@dataclass
class Parsed:
    units: list[Unit] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    parser: str = ""
    _parts: list[str] = field(default_factory=list, repr=False)
    _length: int = field(default=0, repr=False)

    @property
    def raw_text(self) -> str:
        return "".join(self._parts)

    def append(self, text: str, source: dict[str, Any]) -> None:
        if not text:
            return
        if self._parts:
            self._parts.append("\n")
            self._length += 1
        start = self._length
        self._parts.append(text)
        self._length += len(text)
        self.units.append(Unit(start, self._length, source))


def parse_file(path: Path) -> Parsed:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported extension: {suffix}")
    if suffix in {".txt", ".md", ".markdown"}:
        return parse_plain(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pptx":
        return parse_pptx(path)
    return parse_xlsx(path)


def parse_plain(path: Path) -> Parsed:
    text = path.read_text(encoding="utf-8-sig")
    result = Parsed(parser="utf8")
    result.append(text, {"kind": "file"})
    return result


def parse_pdf(path: Path) -> Parsed:
    import pdfplumber

    result = Parsed(parser="pdfplumber")
    with pdfplumber.open(path) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text(layout=False)
            if text:
                result.append(text, {"kind": "pdf_page", "page": number})
            else:
                result.warnings.append(f"Page {number} has no extractable text (scan/image/blank); check OCR manually")
    return result


def parse_docx(path: Path) -> Parsed:
    from docx import Document
    from docx.table import Table
    from zipfile import ZipFile

    result = Parsed(parser="python-docx")
    document = Document(path)
    paragraph_number = 0
    table_number = 0
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            table_number += 1
            for row_number, row in enumerate(block.rows, start=1):
                for column_number, cell in enumerate(row.cells, start=1):
                    if cell.text:
                        result.append(cell.text, {
                            "kind": "docx_cell",
                            "table": table_number,
                            "row": row_number,
                            "column": column_number,
                        })
        else:
            paragraph_number += 1
            if block.text:
                result.append(block.text, {"kind": "docx_paragraph", "paragraph": paragraph_number})
    for section_no, section in enumerate(document.sections, start=1):
        for kind, part in (("header", section.header), ("footer", section.footer)):
            for paragraph_no, paragraph in enumerate(part.paragraphs, start=1):
                if paragraph.text:
                    result.append(paragraph.text, {
                        "kind": f"docx_{kind}",
                        "section": section_no,
                        "paragraph": paragraph_no,
                    })
    with ZipFile(path) as archive:
        word_xml = [name for name in archive.namelist() if name.startswith("word/") and name.endswith(".xml")]
        for name in word_xml:
            if name.startswith("word/embeddings/"):
                continue
            if name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}:
                result.warnings.append(f"{name}: additional content not parsed; inspect original")
            elif name == "word/document.xml":
                xml = archive.read(name)
                if b"txbxContent" in xml or b"w:ins" in xml or b"w:del" in xml:
                    result.warnings.append("DOCX has textboxes or tracked revisions; review original for omissions")
    return result


def parse_pptx(path: Path) -> Parsed:
    """Read slide text, tables, groups and presenter notes without OCR.

    Slide/source locations are retained. Shape enumeration follows document
    XML/z-order rather than claiming a guaranteed visual reading order.
    Embedded images, charts and SmartArt cannot be losslessly converted to
    plain text; report that limitation per slide instead of inventing text.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    result = Parsed(parser="python-pptx")
    deck = Presentation(path)
    for slide_number, slide in enumerate(deck.slides, start=1):
        written = 0
        omitted_images = 0
        omitted_objects = 0

        def read_shapes(shapes, ancestry: tuple[int, ...] = ()) -> None:
            nonlocal written, omitted_images, omitted_objects
            for shape in shapes:
                ids = (*ancestry, shape.shape_id)
                base = {
                    "slide": slide_number,
                    "shape_ids": list(ids),
                    "shape_name": shape.name,
                }
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    read_shapes(shape.shapes, ids)
                elif shape.has_table:
                    for row_number, row in enumerate(shape.table.rows, start=1):
                        for column_number, cell in enumerate(row.cells, start=1):
                            if cell.text:
                                result.append(cell.text, {
                                    **base,
                                    "kind": "pptx_table_cell",
                                    "row": row_number,
                                    "column": column_number,
                                })
                                written += 1
                elif shape.has_text_frame:
                    if shape.text:
                        result.append(shape.text, {**base, "kind": "pptx_text"})
                        written += 1
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    omitted_images += 1
                elif shape.has_chart or shape.element.tag.endswith("}graphicFrame"):
                    omitted_objects += 1

        read_shapes(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                result.append(notes.text, {
                    "kind": "pptx_notes",
                    "slide": slide_number,
                })
        if omitted_images:
            result.warnings.append(
                f"Slide {slide_number}: {omitted_images} image(s) not OCR-processed; text inside images is not indexed"
            )
        if omitted_objects:
            result.warnings.append(
                f"Slide {slide_number}: {omitted_objects} chart/diagram/embedded object(s) not fully parsed"
            )
        if not written:
            result.warnings.append(
                f"Slide {slide_number}: no extractable slide text; check images/diagrams or blank slide"
            )
    return result


def parse_xlsx(path: Path) -> Parsed:
    from openpyxl import load_workbook

    result = Parsed(parser="openpyxl")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for worksheet in workbook.worksheets:
            nonempty = 0
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    nonempty += 1
                    if nonempty > 150_000:
                        raise ValueError(f"Sheet {worksheet.title} exceeds 150000 populated cells")
                    result.append(f"{cell.coordinate} = {value}", {
                        "kind": "xlsx_cell",
                        "sheet": worksheet.title,
                        "cell": cell.coordinate,
                    })
                    if cell.data_type == "f":
                        result.warnings.append(
                            f"Formula at {worksheet.title}!{cell.coordinate}: formula text indexed, cached value not evaluated"
                        )
    finally:
        workbook.close()
    return result
