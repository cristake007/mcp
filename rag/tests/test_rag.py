"""Fidelity and RAG database smoke tests; no MCP imports or model downloads."""
import os
import sys
import tempfile
import unittest
from pathlib import Path

# Run with python -m unittest discover -s rag/tests -v from repository root.
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ragcore.articles import articles, article_by_label, article_for_hit
from ragcore.chunking import Chunk, make_chunks, verify_coverage
from ragcore.parsing import Parsed, parse_file
from ragcore.store import walk_documents


class ParserTests(unittest.TestCase):
    def test_articles_span_across_pages_and_preserve_exact_text(self):
        raw = ("INTRO\nArt. 33 (1) Verificarea se face anual.\n"
               "(2) Evaluarea este practica.\n"
               "15\nMONITORUL OFICIAL\n"
               "Art. 34 (1) Inspectorul completeaza raportul.\n"
               "(2) Documentele se arhiveaza.\nArt. 35 (1) Alte cerinte.")
        parsed = articles(raw)
        self.assertEqual([part.label for part in parsed], ["33", "34", "35"])
        art33 = article_by_label(raw, "Art. 33")
        self.assertIn("MONITORUL OFICIAL", raw[art33.start:art33.end])
        self.assertNotIn("Art. 34", raw[art33.start:art33.end])
        art34 = article_by_label(raw, "34")
        self.assertIn("(2) Documentele se arhiveaza.", raw[art34.start:art34.end])
        self.assertEqual(article_for_hit(raw, art34.start + 5, art34.end - 5), art34)
        self.assertEqual(art34.info(include_text=raw)["text"], raw[art34.start:art34.end])
        with self.assertRaises(ValueError):
            article_by_label(raw + "\nArt. 34 (3) Repetat.", "34")

    def test_exact_chunks_cover_all_whitespace_and_unicode(self):
        document = Parsed(parser="test")
        document.append("Știință   ISO 9001\nArt. 15", {"page": 1})
        document.append("\tObligație: 10 ani.\nFinal.", {"page": 2})
        chunks = make_chunks(document.raw_text, document.units, max_chars=18, overlap=4)
        verify_coverage(document.raw_text, chunks)
        self.assertEqual(chunks[0].text, document.raw_text[chunks[0].start:chunks[0].end])
        self.assertEqual(chunks[-1].end, len(document.raw_text))
        self.assertTrue(any({"page": 2} in chunk.locations for chunk in chunks))
        with self.assertRaises(ValueError):
            verify_coverage(document.raw_text, [
                Chunk(1, len(document.raw_text), document.raw_text[1:], [])
            ])

    def test_recurse_without_following_links(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch) / "documents"
            (root / "deep" / "nested").mkdir(parents=True)
            (root / "deep" / "nested" / "a.txt").write_text("ISO 9001", encoding="utf-8")
            (root / "deep" / "nested" / "skip.png").write_bytes(b"x")
            (root / "external.txt").write_text("secret", encoding="utf-8")
            (root / "outside-link.txt").symlink_to(root / "external.txt")
            found = list(walk_documents(root))
            self.assertEqual(sorted(path for path, _ in found), ["deep/nested/a.txt", "external.txt"])

    def test_word_preserves_tables_paragraphs_headers(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as scratch:
            path = Path(scratch) / "procedure.docx"
            doc = Document()
            doc.add_paragraph("Articolul 15 - păstrare 10 ani")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Parametru"
            table.cell(0, 1).text = "99.25"
            doc.sections[0].header.paragraphs[0].text = "Antet ISO"
            doc.save(path)
            parsed = parse_file(path)
            self.assertIn("Articolul 15", parsed.raw_text)
            self.assertIn("99.25", parsed.raw_text)
            self.assertIn("Antet ISO", parsed.raw_text)
            self.assertTrue(any(x.source.get("kind") == "docx_cell" for x in parsed.units))

    def test_powerpoint_slides_tables_and_speaker_notes(self):
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as scratch:
            path = Path(scratch) / "training.pptx"
            presentation = Presentation()
            first = presentation.slides.add_slide(presentation.slide_layouts[6])
            first.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = (
                "ISO 9001:2015 - Articolul 15"
            )
            table = first.shapes.add_table(
                2, 2, Inches(1), Inches(2), Inches(6), Inches(2)
            ).table
            table.cell(0, 0).text = "Obligatie"
            table.cell(0, 1).text = "10 ani"
            table.cell(1, 0).text = "Termen"
            table.cell(1, 1).text = "24.5"
            first.notes_slide.notes_text_frame.text = "Notite: se pastreaza 10 ani."
            second = presentation.slides.add_slide(presentation.slide_layouts[6])
            second.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = (
                "A doua pagina - ISO 14001"
            )
            presentation.save(path)

            parsed = parse_file(path)
            self.assertEqual(parsed.parser, "python-pptx")
            for value in ["ISO 9001:2015", "10 ani", "24.5", "Notite:", "ISO 14001"]:
                self.assertIn(value, parsed.raw_text)
            sources = [unit.source for unit in parsed.units]
            self.assertTrue(any(x.get("kind") == "pptx_text" and x.get("slide") == 1 for x in sources))
            self.assertTrue(any(x.get("kind") == "pptx_table_cell"
                                and x.get("row") == 1 and x.get("column") == 2
                                and x.get("slide") == 1 for x in sources))
            self.assertTrue(any(x.get("kind") == "pptx_notes" and x.get("slide") == 1
                                for x in sources))
            self.assertTrue(any(x.get("kind") == "pptx_text" and x.get("slide") == 2
                                for x in sources))
            chunks = make_chunks(parsed.raw_text, parsed.units, max_chars=40, overlap=8)
            verify_coverage(parsed.raw_text, chunks)
            self.assertTrue(all(chunk.text == parsed.raw_text[chunk.start:chunk.end]
                                for chunk in chunks))

    def test_excel_preserves_cell_coordinates_and_formulas(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as scratch:
            path = Path(scratch) / "numbers.xlsx"
            book = Workbook()
            ws = book.active
            ws.title = "Registru"
            ws["B2"] = "ISO 14001"
            ws["B3"] = 12.75
            ws["B4"] = "=B3*2"
            book.save(path)
            parsed = parse_file(path)
            self.assertIn("B2 = ISO 14001", parsed.raw_text)
            self.assertIn("B3 = 12.75", parsed.raw_text)
            self.assertIn("B4 = =B3*2", parsed.raw_text)
            self.assertTrue(any(x.source.get("cell") == "B4" for x in parsed.units))
            self.assertTrue(any("Formula" in warning for warning in parsed.warnings))

    def test_pdf_has_page_provenance(self):
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as scratch:
            path = Path(scratch) / "article.pdf"
            pdf = canvas.Canvas(str(path))
            pdf.drawString(72, 720, "Articolul 15 termen 10 ani")
            pdf.showPage()
            pdf.save()
            parsed = parse_file(path)
            self.assertIn("Articolul 15", parsed.raw_text)
            self.assertTrue(any(x.source.get("page") == 1 for x in parsed.units))


@unittest.skipUnless(os.environ.get("RAG_TEST_DB") == "1", "Only with isolated CI PostgreSQL")
class ApiTests(unittest.IsolatedAsyncioTestCase):
    async def test_optional_read_only_http_api(self):
        import httpx
        from ragcore import store
        from ragcore.api import app

        with tempfile.TemporaryDirectory() as scratch:
            file = Path(scratch) / "legal.txt"
            file.write_text(
                "Art. 33 (1) Verificari tehnice.\n(2) Se face anual.\n"
                "Art. 34 (1) ISCIR verifica dosarul.\n(2) Dosarul se semneaza.",
                encoding="utf-8",
            )
            with store.connect() as conn:
                store.init_schema(conn)
                with conn.cursor() as cur:
                    cur.execute("DELETE FROM rag_documents")
                self.assertEqual(store.scan(conn, Path(scratch), use_vectors=False)["updated"], 1)
            prior = os.environ.get("RAG_API_TOKEN")
            os.environ["RAG_API_TOKEN"] = "unit-test-example-unpredictable-token-0123456789abcdef"
            try:
                async with httpx.AsyncClient(
                    transport=httpx.ASGITransport(app=app), base_url="http://rag-api"
                ) as client:
                    denied = await client.post("/search", json={"query": "verificari"})
                    self.assertEqual(denied.status_code, 401)
                    headers = {"Authorization": "Bearer " + os.environ["RAG_API_TOKEN"]}
                    results = await client.post(
                        "/search", json={"query": "verificari"}, headers=headers
                    )
                    self.assertEqual(results.status_code, 200, results.text)
                    item = results.json()["results"][0]
                    self.assertEqual(item["path"], "legal.txt")
                    self.assertEqual(item["article"]["label"], "33")
                    self.assertIn("(2) Se face anual", item["article"]["text"])
                    text = await client.post(
                        "/text",
                        json={"document_id": item["document_id"], "article": "Art. 33"},
                        headers=headers,
                    )
                    self.assertEqual(text.status_code, 200, text.text)
                    self.assertEqual(text.json()["text"], item["article"]["text"])
                    self.assertTrue(text.json()["exact"])
                    bad_range = await client.post(
                        "/text",
                        json={"document_id": item["document_id"], "start": -1},
                        headers=headers,
                    )
                    self.assertEqual(bad_range.status_code, 400)
                    listing = await client.post(
                        "/documents", json={"offset": 0, "limit": 10}, headers=headers,
                    )
                    self.assertEqual(listing.status_code, 200, listing.text)
                    docs = listing.json()
                    self.assertEqual(docs["total_documents"], 1)
                    self.assertEqual(docs["documents"][0]["document_id"], item["document_id"])
                    self.assertEqual(docs["documents"][0]["warnings_count"], 0)
                    self.assertEqual(docs["source_check"], "not_performed")
                    self.assertIsNone(docs["next_offset"])
                    status = await client.post(
                        "/document-status", json={"document_id": item["document_id"]},
                        headers=headers,
                    )
                    self.assertEqual(status.status_code, 200, status.text)
                    self.assertEqual(status.json()["status"], "indexed_complete")
                    self.assertTrue(status.json()["stored_text_coverage_complete"])
                    self.assertGreater(status.json()["chunks"], 0)
                    self.assertEqual(status.json()["source_check"], "not_performed")
                    denied_list = await client.post("/documents", json={})
                    self.assertEqual(denied_list.status_code, 401)
                    invalid_list = await client.post(
                        "/documents", json={"offset": False}, headers=headers,
                    )
                    self.assertEqual(invalid_list.status_code, 400)
                    invalid_status = await client.post(
                        "/document-status", json={"document_id": 0}, headers=headers,
                    )
                    self.assertEqual(invalid_status.status_code, 400)
                    missing = await client.post(
                        "/document-status", json={"document_id": 9999999}, headers=headers,
                    )
                    self.assertEqual(missing.status_code, 400)
                    from unittest.mock import patch
                    with patch.dict(os.environ, {"RAG_DOCUMENTS": str(scratch)}):
                        corpus = await client.post(
                            "/corpus-health",
                            json={"offset": 0, "limit": 5},
                            headers=headers,
                        )
                        self.assertEqual(corpus.status_code, 200, corpus.text)
                        self.assertEqual(corpus.json()["status"], "healthy")
                        self.assertEqual(corpus.json()["summary"]["current_files"], 1)
                        self.assertEqual(corpus.json()["summary"]["unindexed_files"], 0)
                        file.write_text("Text schimbat între două scanări", encoding="utf-8")
                        changed = await client.post(
                            "/corpus-health", json={}, headers=headers,
                        )
                        self.assertEqual(changed.status_code, 200, changed.text)
                        self.assertEqual(changed.json()["summary"]["modified_files"], 1)
                        invalid = await client.post(
                            "/corpus-health", json={"limit": False}, headers=headers,
                        )
                        self.assertEqual(invalid.status_code, 400)
                    no_auth = await client.post("/corpus-health", json={})
                    self.assertEqual(no_auth.status_code, 401)
            finally:
                if prior is None:
                    os.environ.pop("RAG_API_TOKEN", None)
                else:
                    os.environ["RAG_API_TOKEN"] = prior


@unittest.skipUnless(os.environ.get("RAG_TEST_DB") == "1", "Only with isolated CI PostgreSQL")
class CorpusHealthTests(unittest.TestCase):
    def test_full_inventory_and_mount_safety(self):
        from ragcore import store
        from ragcore.health import corpus_health

        with tempfile.TemporaryDirectory() as scratch:
            folder = Path(scratch)
            (folder / "nested").mkdir()
            (folder / "nested" / "indexed.txt").write_text("Prima versiune", encoding="utf-8")
            missing = folder / "missing.md"
            missing.write_text("Un document care dispare", encoding="utf-8")
            with store.connect() as conn:
                store.init_schema(conn)
                with conn.cursor() as cur:
                    cur.execute("DELETE FROM rag_documents")
                self.assertEqual(store.scan(conn, folder, use_vectors=False)["updated"], 2)
                first = corpus_health(conn, folder)
                self.assertEqual(first["status"], "healthy")
                self.assertEqual(first["summary"]["current_files"], 2)
                self.assertEqual(first["summary"]["indexed_documents"], 2)
                self.assertEqual(first["summary"]["unindexed_files"], 0)
                self.assertTrue(first["source_directory_checked"])
                self.assertTrue(first["read_only"])

                (folder / "nested" / "indexed.txt").write_text(
                    "Document modificat", encoding="utf-8",
                )
                missing.unlink()
                (folder / "new.pdf").write_bytes(b"%PDF-1.5 fake test")
                (folder / "readme.bin").write_bytes(b"Not a supported extension")
                large = folder / "very-large.txt"
                with large.open("wb") as stream:
                    stream.truncate(31 * 1024 * 1024)
                (folder / "symlink.txt").symlink_to(folder / "nested" / "indexed.txt")

                report = corpus_health(conn, folder, limit=2)
                self.assertEqual(report["status"], "attention_required")
                self.assertEqual(report["summary"]["indexed_documents"], 2)
                self.assertEqual(report["summary"]["modified_files"], 1)
                self.assertEqual(report["summary"]["new_files"], 1)
                self.assertEqual(report["summary"]["unindexed_files"], 1)
                self.assertEqual(report["summary"]["missing_files"], 1)
                self.assertEqual(report["summary"]["oversized_files"], 1)
                self.assertEqual(report["summary"]["unsupported_files"], 1)
                self.assertEqual(report["summary"]["ignored_symlinks"], 1)
                self.assertIsNotNone(report["next_offset"])
                all_files = []
                offset = 0
                while True:
                    batch = corpus_health(conn, folder, offset=offset, limit=2)
                    self.assertEqual(batch["summary"], report["summary"])
                    all_files.extend(batch["files"])
                    if batch["next_offset"] is None:
                        break
                    offset = batch["next_offset"]
                self.assertEqual(
                    {entry["path"]: entry["status"] for entry in all_files},
                    {"nested/indexed.txt": "modified", "missing.md": "missing",
                     "new.pdf": "new_unindexed", "readme.bin": "unsupported",
                     "very-large.txt": "oversized", "symlink.txt": "ignored_symlink"},
                )

                for path in (folder / "nested" / "indexed.txt", folder / "new.pdf",
                             folder / "readme.bin", large, folder / "symlink.txt"):
                    path.unlink()
                empty = corpus_health(conn, folder)
                self.assertTrue(empty["source_directory_empty"])
                self.assertEqual(empty["status"], "source_directory_empty_check_mount")
                self.assertIsNotNone(empty["source_mount_warning"])
                self.assertEqual(empty["summary"]["missing_files"], 2)

                with self.assertRaisesRegex(RuntimeError, "missing"):
                    corpus_health(conn, folder / "not-mounted")

    def test_validation_and_empty_index(self):
        from ragcore import store
        from ragcore.health import corpus_health

        with tempfile.TemporaryDirectory() as scratch:
            folder = Path(scratch)
            with store.connect() as conn:
                store.init_schema(conn)
                with conn.cursor() as cur:
                    cur.execute("DELETE FROM rag_documents")
                report = corpus_health(conn, folder)
                self.assertEqual(report["status"], "healthy")
                self.assertEqual(report["summary"]["indexed_documents"], 0)
                self.assertEqual(report["summary"]["supported_files_on_disk"], 0)
                for offset, limit in ((False, 10), (0, True), (-1, 10), (0, 0)):
                    with self.assertRaises(ValueError):
                        corpus_health(conn, folder, offset=offset, limit=limit)


@unittest.skipUnless(os.environ.get("RAG_TEST_DB") == "1", "Only with isolated CI PostgreSQL")
class StoreTests(unittest.TestCase):
    def test_incremental_search_prune_and_exact_retrieval(self):
        from ragcore import store

        with tempfile.TemporaryDirectory() as scratch:
            folder = Path(scratch)
            (folder / "nested").mkdir()
            source = folder / "nested" / "procedure.txt"
            source.write_text("Articolul 15 spune: 10 ani.\nISO 9001", encoding="utf-8")
            second = folder / "other.md"
            second.write_text("Instruire ISO 14001", encoding="utf-8")
            with store.connect() as conn:
                store.init_schema(conn)
                with conn.cursor() as cur:
                    cur.execute("DELETE FROM rag_documents")
                first = store.scan(conn, folder, use_vectors=False)
                self.assertEqual(first["updated"], 2)
                same = store.scan(conn, folder, use_vectors=False)
                self.assertEqual(same["unchanged"], 2)

                # A natural-language question must have a lexical fallback even
                # if strict Romanian FTS matches no entire phrase.
                natural = store.search(
                    conn, "Care sunt cerintele pentru instruire ISO 14001?",
                    use_vectors=False,
                )
                self.assertTrue(natural)
                self.assertEqual(natural[0]["path"], "other.md")
                results = store.search(conn, "ISO 9001", use_vectors=False)
                self.assertTrue(results, "Hybrid lexical search should return candidates")
                self.assertEqual(results[0]["path"], "nested/procedure.txt")
                row = results[0]
                _, literal = store.get_text(
                    conn, row["document_id"], row["char_start"], row["char_end"]
                )
                self.assertEqual(literal, row["content"])
                self.assertEqual(store.verify_index(conn)["content_mismatches"], 0)
                self.assertEqual(store.verify_index(conn)["coverage_errors"], [])
                listing = store.list_documents(conn, offset=0, limit=1)
                self.assertEqual(listing["total_documents"], 2)
                self.assertEqual(len(listing["documents"]), 1)
                self.assertEqual(listing["next_offset"], 1)
                self.assertEqual(store.list_documents(conn, offset=1, limit=1)["next_offset"], None)
                status = store.document_status(conn, row["document_id"])
                self.assertEqual(status["status"], "indexed_complete")
                self.assertFalse(status["integrity_errors"])
                with conn.cursor() as cur:
                    cur.execute(
                        "UPDATE rag_chunks SET content='corrupted' "
                        "WHERE document_id=%s AND ordinal=0",
                        (row["document_id"],),
                    )
                broken = store.document_status(conn, row["document_id"])
                self.assertEqual(broken["status"], "incomplete")
                self.assertFalse(broken["stored_text_coverage_complete"])
                with conn.cursor() as cur:
                    cur.execute(
                        "UPDATE rag_chunks SET content=substring(d.raw_text "
                        "FROM rag_chunks.char_start+1 FOR rag_chunks.char_end-rag_chunks.char_start) "
                        "FROM rag_documents d WHERE rag_chunks.document_id=d.id "
                        "AND rag_chunks.document_id=%s AND rag_chunks.ordinal=0",
                        (row["document_id"],),
                    )
                self.assertEqual(
                    store.document_status(conn, row["document_id"])["status"], "indexed_complete",
                )

                source.write_text("Articolul 15: 25 de ani!\nISO 9001", encoding="utf-8")
                changed = store.scan(conn, folder, use_vectors=False)
                self.assertEqual(changed["updated"], 1)
                source.unlink()
                removed = store.scan(conn, folder, use_vectors=False, prune=True)
                self.assertEqual(removed["deleted"], 1)
                self.assertFalse(store.search(conn, "ISO 9001", use_vectors=False))


if __name__ == "__main__":
    unittest.main()
